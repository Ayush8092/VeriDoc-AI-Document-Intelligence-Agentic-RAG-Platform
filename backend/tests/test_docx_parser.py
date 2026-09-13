from pathlib import Path

import docx
import pytest

from app.documents.models import BlockType
from app.documents.parsers.docx_parser import parse_docx


def _build_docx(path: Path) -> None:
    d = docx.Document()
    d.add_heading("Doc Title", level=1)
    d.add_paragraph("Intro paragraph.")
    d.add_heading("Section A", level=2)
    d.add_paragraph("Section A body.")
    table = d.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "Item", "Qty"
    row = table.add_row().cells
    row[0].text, row[1].text = "Widget", "4"
    d.save(str(path))


def test_parse_docx_extracts_headings_paragraphs_and_table(tmp_path: Path):
    f = tmp_path / "doc.docx"
    _build_docx(f)

    doc = parse_docx(f)

    headings = [b.text for b in doc.blocks if b.block_type == BlockType.HEADING]
    assert "Doc Title" in headings
    assert "Section A" in headings

    tables = [b for b in doc.blocks if b.block_type == BlockType.TABLE]
    assert len(tables) == 1
    assert tables[0].table.rows == (("Item", "Qty"), ("Widget", "4"))


def test_parse_docx_missing_file_raises(tmp_path: Path):
    with pytest.raises(ValueError):
        parse_docx(tmp_path / "does_not_exist.docx")


def test_parse_docx_empty_document_raises(tmp_path: Path):
    f = tmp_path / "empty.docx"
    docx.Document().save(str(f))
    with pytest.raises(ValueError):
        parse_docx(f)


def test_docx_paragraph_image_is_extracted(tmp_path):
    from docx import Document
    from PIL import Image

    image_path = tmp_path / "figure.png"

    image = Image.new("RGB", (300, 200), "white")
    image.save(image_path)

    doc = Document()

    doc.add_paragraph("Before image")

    image_paragraph = doc.add_paragraph()
    image_paragraph.add_run().add_picture(str(image_path))

    doc.add_paragraph("After image")

    docx_path = tmp_path / "sample.docx"
    doc.save(docx_path)

    extracted = parse_docx(docx_path)

    visual_blocks = [
        block
        for block in extracted.blocks
        if block.visual is not None
    ]

    assert len(visual_blocks) == 1

    visual = visual_blocks[0].visual

    assert visual is not None
    assert visual.page_number == 1
    assert visual.bbox is None
    assert visual.extraction_method == "docx_inline_shape"
    assert visual.metadata["width"] == 300
    assert visual.metadata["height"] == 200
def test_docx_paragraph_multiple_images(tmp_path):
    from docx import Document
    from PIL import Image

    image_path_1 = tmp_path / "figure1.png"
    image_path_2 = tmp_path / "figure2.png"

    Image.new("RGB", (300, 200), "white").save(image_path_1)
    Image.new("RGB", (400, 250), "white").save(image_path_2)

    doc = Document()

    paragraph = doc.add_paragraph()
    paragraph.add_run().add_picture(str(image_path_1))
    paragraph.add_run().add_picture(str(image_path_2))

    docx_path = tmp_path / "multi.docx"
    doc.save(docx_path)

    extracted = parse_docx(docx_path)

    visuals = [
        block
        for block in extracted.blocks
        if block.visual is not None
    ]

    assert len(visuals) == 2

    assert len(visuals) == 2

    assert sorted(
        (
            block.visual.metadata["width"],
            block.visual.metadata["height"],
        )
        for block in visuals
    ) == [
        (300, 200),
        (400, 250),
    ]